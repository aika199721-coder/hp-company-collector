"""Phase 8 Step 3 Excel exporter tests."""

from openpyxl import load_workbook

from export.excel import SHEET_NAMES, ExcelExporter
from export.models import ExportSnapshot
from export.repository import ExportRepository


def test_excel_creates_all_sheets_links_and_text_formats(populated_database, tmp_path) -> None:
    snapshot = ExportRepository(populated_database).read_snapshot()
    path = tmp_path / "companies.xlsx"

    ExcelExporter().export(path, snapshot)
    workbook = load_workbook(path)

    assert workbook.sheetnames == list(SHEET_NAMES)
    assert "検索履歴" in workbook.sheetnames
    sheet = workbook["営業対象"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    url_cell = sheet.cell(2, headers["公式HP"])
    assert url_cell.hyperlink is not None
    assert sheet.cell(2, headers["電話番号"]).number_format == "@"
    assert sheet.cell(2, headers["郵便番号"]).number_format == "@"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None


def test_empty_snapshot_still_has_sheets_and_headers(tmp_path) -> None:
    path = tmp_path / "empty.xlsx"

    ExcelExporter().export(path, ExportSnapshot((), ()))
    workbook = load_workbook(path)

    assert workbook.sheetnames == list(SHEET_NAMES)
    assert "検索履歴" in workbook.sheetnames
    assert workbook["全件"].max_row == 1


def test_formula_like_values_are_escaped(populated_database, tmp_path) -> None:
    snapshot = ExportRepository(populated_database).read_snapshot()
    altered = snapshot.with_first_company_name("=HYPERLINK(\"bad\")")
    path = tmp_path / "safe.xlsx"

    ExcelExporter().export(path, altered)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["全件"]
    headers = {cell.value: cell.column for cell in sheet[1]}

    assert str(sheet.cell(2, headers["法人名"]).value).startswith("'")
