"""Offline live report, review history, and accuracy integration tests."""

import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

from live.report import LIVE_SHEETS, LiveReportService
from live.review import REVIEW_HEADERS, ManualReviewService


def _review(path: Path, url: str, judgment: str = "OK") -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(REVIEW_HEADERS)
        writer.writerow([url, *(judgment for _ in range(8)), "確認メモ"])


def test_live_report_has_all_sheets_and_generates_review_csv(
    populated_database, tmp_path: Path
) -> None:
    workbook = tmp_path / "live_validation_report.xlsx"
    review = tmp_path / "live_review.csv"

    LiveReportService(populated_database).export(workbook, review)

    assert load_workbook(workbook).sheetnames == list(LIVE_SHEETS)
    assert review.read_bytes().startswith(b"\xef\xbb\xbf")
    assert "未確認" in review.read_text(encoding="utf-8-sig")


def test_review_import_preserves_history_and_automatic_scoring(
    populated_database, tmp_path: Path
) -> None:
    path = tmp_path / "review.csv"
    _review(path, "https://official.example/")
    service = ManualReviewService(populated_database)
    with populated_database.connect() as connection:
        before = connection.execute("SELECT is_business_target FROM scoring_results").fetchone()[0]

    assert service.import_csv(path, "reviewer-a") == 1
    assert service.import_csv(path, "reviewer-b") == 1
    report = service.accuracy()

    with populated_database.connect() as connection:
        review_count = connection.execute("SELECT COUNT(*) FROM manual_reviews").fetchone()[0]
        after = connection.execute("SELECT is_business_target FROM scoring_results").fetchone()[0]
    assert review_count == 2
    assert before == after
    assert report.company_name_accuracy == 1.0
    assert report.statistically_sufficient is False


def test_unknown_review_value_has_row_number(populated_database, tmp_path: Path) -> None:
    path = tmp_path / "review.csv"
    _review(path, "https://official.example/", "不明")

    with pytest.raises(ValueError, match="row 2"):
        ManualReviewService(populated_database).import_csv(path)
